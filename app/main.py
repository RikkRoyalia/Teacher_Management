import json
import os
import secrets
from contextlib import asynccontextmanager
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from typing import Annotated, Any

from fastapi import (
    Depends,
    FastAPI,
    File,
    Form,
    HTTPException,
    Query,
    Request,
    UploadFile,
    status,
)
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from sqlalchemy import func, or_, select, update
from sqlalchemy.exc import IntegrityError, SQLAlchemyError
from sqlalchemy.orm import Session

from app.audit import record_audit
from app.database import DATABASE_URL, create_database, migrate_database
from app.models import (
    AppSetting,
    AuditLog,
    ChangeRequest,
    ChangeRequestEvent,
    ChecklistItem,
    Department,
    EmploymentContract,
    EvidenceDocument,
    LoginSession,
    Qualification,
    Teacher,
    TeacherChecklistRequirement,
    User,
    WorkHistory,
)
from app.schemas import (
    ChangeDecision,
    ChangeRequestCreate,
    ChangeRequestRead,
    ChangeRequestUpdate,
    ChecklistItemCreate,
    ChecklistItemRead,
    ChecklistItemUpdate,
    ContractCreate,
    ContractRead,
    ContractUpdate,
    DepartmentCreate,
    DepartmentRead,
    DepartmentUpdate,
    EvidenceDocumentRead,
    LoginInput,
    PasswordChange,
    QualificationCreate,
    QualificationRead,
    QualificationUpdate,
    TeacherCreate,
    TeacherChecklistRequirementRead,
    TeacherChecklistRequirementUpdate,
    TeacherRead,
    TeacherUpdate,
    UserCreate,
    UserRead,
    UserUpdate,
    WorkHistoryCreate,
    WorkHistoryRead,
)
from app.security import (
    CSRF_COOKIE,
    SESSION_COOKIE,
    DbSession,
    get_current_user,
    hash_password,
    hash_secret,
    issue_session,
    require_roles,
    verify_password,
)

STATIC_DIR = Path(__file__).parent / "static"
UPLOAD_DIR = Path(
    os.getenv("TEACHER_MANAGEMENT_UPLOAD_DIR", str(Path(__file__).parent.parent / "private_uploads"))
).resolve()
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
ALLOWED_UPLOADS = {
    ".pdf": ("application/pdf", lambda data: data.startswith(b"%PDF-")),
    ".png": ("image/png", lambda data: data.startswith(b"\x89PNG\r\n\x1a\n")),
    ".jpg": ("image/jpeg", lambda data: data.startswith(b"\xff\xd8\xff")),
    ".jpeg": ("image/jpeg", lambda data: data.startswith(b"\xff\xd8\xff")),
}
IMPORTANT_TEACHER_FIELDS = {
    "department_id",
    "academic_degree",
    "academic_title",
    "employment_status",
}
APPROVABLE_FIELDS = IMPORTANT_TEACHER_FIELDS
HR_USER = Annotated[User, Depends(require_roles("hr"))]
MANAGER_USER = Annotated[User, Depends(require_roles("manager"))]
ADMIN_USER = Annotated[User, Depends(require_roles("admin"))]
HR_OR_MANAGER = Annotated[User, Depends(require_roles("hr", "manager"))]


def create_app(database_url: str = DATABASE_URL) -> FastAPI:
    engine, session_factory = create_database(database_url)

    @asynccontextmanager
    async def lifespan(_: FastAPI):
        migrate_database(database_url, engine)
        UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
        yield
        engine.dispose()

    application = FastAPI(
        title="Teacher Management API",
        description="Hệ thống quản lý hồ sơ giảng viên, quy trình và báo cáo nội bộ.",
        version="1.0.0",
        lifespan=lifespan,
    )
    application.state.session_factory = session_factory
    application.state.upload_dir = UPLOAD_DIR
    application.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")

    @application.middleware("http")
    async def security_headers(request: Request, call_next):
        response = await call_next(request)
        response.headers.setdefault("Cache-Control", "no-store")
        response.headers["Content-Security-Policy"] = (
            "default-src 'self'; script-src 'self'; style-src 'self'; "
            "img-src 'self' data:; connect-src 'self'; object-src 'none'; "
            "base-uri 'self'; form-action 'self'; frame-ancestors 'none'"
        )
        response.headers["Referrer-Policy"] = "no-referrer"
        response.headers["X-Content-Type-Options"] = "nosniff"
        response.headers["X-Frame-Options"] = "DENY"
        return response

    @application.get("/", include_in_schema=False)
    def frontend() -> FileResponse:
        return FileResponse(STATIC_DIR / "index.html")

    @application.get("/health", tags=["system"])
    def health_check() -> dict[str, str]:
        try:
            with session_factory() as session:
                session.execute(select(1))
        except SQLAlchemyError as error:
            raise HTTPException(status_code=503, detail="Database is unavailable.") from error
        return {"status": "ok"}

    @application.post("/auth/login", tags=["authentication"])
    def login(payload: LoginInput, request: Request, db: DbSession) -> dict[str, Any]:
        normalized_username = payload.username.strip().casefold()
        user = db.scalar(
            select(User).where(func.lower(User.username) == normalized_username)
        )
        now = datetime.now(timezone.utc)
        if user is None:
            record_audit(db, None, "login_failed", "user", None, "failure")
            db.commit()
            raise HTTPException(
                status_code=status.HTTP_401_UNAUTHORIZED,
                detail="Invalid username or password.",
            )
        if user.locked_until and _as_utc(user.locked_until) > now:
            record_audit(db, None, "login_blocked", "user", user.id, "failure")
            db.commit()
            raise HTTPException(
                status_code=status.HTTP_429_TOO_MANY_REQUESTS,
                detail="Account is temporarily locked. Try again later.",
            )
        if not user.is_active or not verify_password(payload.password, user.password_hash):
            user.failed_login_attempts += 1
            if user.failed_login_attempts >= 5:
                user.failed_login_attempts = 0
                user.locked_until = now + timedelta(minutes=15)
            record_audit(db, None, "login_failed", "user", user.id, "failure")
            db.commit()
            raise HTTPException(
                status_code=status.HTTP_401_UNAUTHORIZED,
                detail="Invalid username or password.",
            )

        user.failed_login_attempts = 0
        user.locked_until = None
        token, csrf_token = issue_session(db, user)
        record_audit(db, user, "login", "user", user.id)
        db.commit()
        response_user = UserRead.model_validate(user).model_dump(mode="json")
        response_user["csrf_token"] = csrf_token
        from fastapi.responses import JSONResponse

        response = JSONResponse(response_user)
        response.set_cookie(
            SESSION_COOKIE,
            token,
            httponly=True,
            secure=False,
            samesite="strict",
            max_age=8 * 60 * 60,
            path="/",
        )
        response.set_cookie(
            CSRF_COOKIE,
            csrf_token,
            httponly=False,
            secure=False,
            samesite="strict",
            max_age=8 * 60 * 60,
            path="/",
        )
        return response

    @application.get("/auth/session", tags=["authentication"])
    def auth_session(
        request: Request,
        db: DbSession,
        user: Annotated[User, Depends(get_current_user)],
    ) -> dict[str, Any]:
        login_session: LoginSession = request.state.login_session
        csrf_token = request.cookies.get(CSRF_COOKIE, "")
        if not csrf_token or not secrets.compare_digest(
            hash_secret(csrf_token), login_session.csrf_token_hash
        ):
            csrf_token = secrets.token_urlsafe(32)
            login_session.csrf_token_hash = hash_secret(csrf_token)
            db.commit()
        result = UserRead.model_validate(user).model_dump(mode="json")
        result["csrf_token"] = csrf_token
        from fastapi.responses import JSONResponse

        response = JSONResponse(result)
        response.set_cookie(
            CSRF_COOKIE,
            csrf_token,
            httponly=False,
            secure=False,
            samesite="strict",
            max_age=8 * 60 * 60,
            path="/",
        )
        return response

    @application.post("/auth/logout", tags=["authentication"])
    def logout(
        request: Request,
        db: DbSession,
        user: Annotated[User, Depends(require_roles("admin", "hr", "manager"))],
    ) -> dict[str, str]:
        login_session: LoginSession = request.state.login_session
        db.delete(login_session)
        record_audit(db, user, "logout", "user", user.id)
        db.commit()
        from fastapi.responses import JSONResponse

        response = JSONResponse({"status": "ok"})
        response.delete_cookie(SESSION_COOKIE, path="/", samesite="strict")
        response.delete_cookie(CSRF_COOKIE, path="/", samesite="strict")
        return response

    @application.post("/auth/password", tags=["authentication"])
    def change_password(
        payload: PasswordChange,
        request: Request,
        db: DbSession,
        user: Annotated[User, Depends(require_roles("admin", "hr", "manager"))],
    ) -> dict[str, str]:
        _require_csrf(request)
        if not verify_password(payload.current_password, user.password_hash):
            raise HTTPException(status_code=400, detail="Current password is incorrect.")
        if payload.current_password == payload.new_password:
            raise HTTPException(
                status_code=400, detail="New password must differ from the current password."
            )
        user.password_hash = hash_password(payload.new_password)
        for existing in db.scalars(
            select(LoginSession).where(LoginSession.user_id == user.id)
        ):
            db.delete(existing)
        record_audit(db, user, "password_changed", "user", user.id)
        db.commit()
        from fastapi.responses import JSONResponse

        response = JSONResponse({"status": "password_changed"})
        response.delete_cookie(SESSION_COOKIE, path="/", samesite="strict")
        response.delete_cookie(CSRF_COOKIE, path="/", samesite="strict")
        return response

    @application.get("/users", response_model=list[UserRead], tags=["accounts"])
    def list_users(db: DbSession, _: ADMIN_USER) -> list[User]:
        return list(db.scalars(select(User).order_by(User.username)))

    @application.post(
        "/users",
        response_model=UserRead,
        status_code=status.HTTP_201_CREATED,
        tags=["accounts"],
    )
    def create_user(payload: UserCreate, db: DbSession, actor: ADMIN_USER) -> User:
        user = User(
            username=payload.username.lower(),
            full_name=payload.full_name,
            password_hash=hash_password(payload.password),
            role=payload.role,
        )
        db.add(user)
        try:
            db.flush()
            record_audit(
                db, actor, "user_created", "user", user.id, details={"role": user.role}
            )
            db.commit()
        except IntegrityError as error:
            db.rollback()
            raise HTTPException(status_code=409, detail="Username already exists.") from error
        db.refresh(user)
        return user

    @application.patch("/users/{user_id}", response_model=UserRead, tags=["accounts"])
    def update_user(
        user_id: int,
        payload: UserUpdate,
        db: DbSession,
        actor: ADMIN_USER,
    ) -> User:
        target = db.get(User, user_id)
        if target is None:
            raise HTTPException(status_code=404, detail="User not found.")
        changes = payload.model_dump(exclude_unset=True)
        if target.id == actor.id and (
            changes.get("is_active") is False
            or ("role" in changes and changes["role"] != actor.role)
        ):
            raise HTTPException(status_code=400, detail="You cannot disable or demote yourself.")
        if target.role == "admin" and (
            changes.get("is_active") is False or changes.get("role") not in {None, "admin"}
        ):
            active_admins = db.scalar(
                select(func.count()).select_from(User).where(
                    User.role == "admin", User.is_active.is_(True)
                )
            )
            if active_admins <= 1:
                raise HTTPException(status_code=409, detail="At least one active administrator is required.")
        if (
            ("is_active" in changes and changes["is_active"] != target.is_active)
            or ("role" in changes and changes["role"] != target.role)
        ):
            for login_session in db.scalars(
                select(LoginSession).where(LoginSession.user_id == target.id)
            ):
                db.delete(login_session)
        for field, value in changes.items():
            setattr(target, field, value)
        record_audit(db, actor, "user_updated", "user", target.id, details=changes)
        db.commit()
        db.refresh(target)
        return target

    @application.post(
        "/departments",
        response_model=DepartmentRead,
        status_code=status.HTTP_201_CREATED,
        tags=["departments"],
    )
    def create_department(payload: DepartmentCreate, db: DbSession, actor: HR_USER) -> Department:
        _validate_parent_department(db, payload.parent_id)
        department = Department(**payload.model_dump())
        db.add(department)
        try:
            db.flush()
            record_audit(db, actor, "department_created", "department", department.id)
            db.commit()
        except IntegrityError as error:
            db.rollback()
            raise HTTPException(status_code=409, detail="Department code already exists.") from error
        db.refresh(department)
        return department

    @application.get(
        "/departments",
        response_model=list[DepartmentRead],
        tags=["departments"],
    )
    def list_departments(
        db: DbSession,
        _: Annotated[User, Depends(require_roles("hr", "manager"))],
        include_inactive: bool = False,
    ) -> list[Department]:
        statement = select(Department).order_by(Department.name, Department.id)
        if not include_inactive:
            statement = statement.where(Department.is_active.is_(True))
        return list(db.scalars(statement))

    @application.patch(
        "/departments/{department_id}",
        response_model=DepartmentRead,
        tags=["departments"],
    )
    def update_department(
        department_id: int,
        payload: DepartmentUpdate,
        db: DbSession,
        actor: HR_USER,
    ) -> Department:
        department = db.get(Department, department_id)
        if department is None:
            raise HTTPException(status_code=404, detail="Department not found.")
        changes = payload.model_dump(exclude_unset=True)
        if "parent_id" in changes:
            _validate_parent_department(db, changes["parent_id"], department_id)
        for field, value in changes.items():
            setattr(department, field, value)
        try:
            record_audit(db, actor, "department_updated", "department", department.id)
            db.commit()
        except IntegrityError as error:
            db.rollback()
            raise HTTPException(status_code=409, detail="Department code already exists.") from error
        db.refresh(department)
        return department

    @application.post(
        "/teachers",
        response_model=TeacherRead,
        status_code=status.HTTP_201_CREATED,
        tags=["teachers"],
    )
    def create_teacher(payload: TeacherCreate, db: DbSession, actor: HR_USER) -> Teacher:
        department = db.get(Department, payload.department_id)
        if department is None or not department.is_active:
            raise HTTPException(status_code=404, detail="Active department not found.")
        teacher = Teacher(**payload.model_dump())
        db.add(teacher)
        try:
            db.flush()
            for item in db.scalars(
                select(ChecklistItem).where(ChecklistItem.is_active.is_(True))
            ):
                db.add(
                    TeacherChecklistRequirement(
                        teacher_id=teacher.id,
                        checklist_item_id=item.id,
                        status=_default_checklist_status(item),
                    )
                )
            db.add(
                WorkHistory(
                    teacher_id=teacher.id,
                    department_id=teacher.department_id,
                    position_title=teacher.academic_title,
                    starts_on=date.today(),
                    reason="Initial employment record",
                )
            )
            record_audit(db, actor, "teacher_created", "teacher", teacher.id)
            db.commit()
        except IntegrityError as error:
            db.rollback()
            raise HTTPException(status_code=409, detail="Teacher employee code already exists.") from error
        db.refresh(teacher)
        return teacher

    @application.get("/teachers", response_model=list[TeacherRead], tags=["teachers"])
    def list_teachers(
        db: DbSession,
        _: HR_USER,
        q: Annotated[str | None, Query(min_length=1, max_length=160)] = None,
        department_id: int | None = Query(default=None, gt=0),
        academic_degree: Annotated[str | None, Query(min_length=1, max_length=80)] = None,
        employment_status: str | None = Query(
            default=None, pattern="^(active|paused|ended)$"
        ),
    ) -> list[Teacher]:
        statement = select(Teacher)
        if q:
            search_term = f"%{q}%"
            statement = statement.where(
                or_(
                    Teacher.employee_code.ilike(search_term),
                    Teacher.full_name.ilike(search_term),
                )
            )
        if department_id is not None:
            statement = statement.where(Teacher.department_id == department_id)
        if academic_degree:
            statement = statement.where(Teacher.academic_degree == academic_degree)
        if employment_status:
            statement = statement.where(Teacher.employment_status == employment_status)
        return list(db.scalars(statement.order_by(Teacher.full_name, Teacher.id)))

    @application.get(
        "/teachers/{teacher_id}",
        response_model=TeacherRead,
        tags=["teachers"],
    )
    def get_teacher(teacher_id: int, db: DbSession, _: HR_USER) -> Teacher:
        teacher = db.get(Teacher, teacher_id)
        if teacher is None:
            raise HTTPException(status_code=404, detail="Teacher not found.")
        return teacher

    @application.patch(
        "/teachers/{teacher_id}",
        response_model=TeacherRead,
        tags=["teachers"],
    )
    def update_teacher(
        teacher_id: int,
        payload: TeacherUpdate,
        db: DbSession,
        actor: HR_USER,
    ) -> Teacher:
        teacher = db.get(Teacher, teacher_id)
        if teacher is None:
            raise HTTPException(status_code=404, detail="Teacher not found.")
        changes = payload.model_dump(exclude_unset=True)
        approval_fields = IMPORTANT_TEACHER_FIELDS.intersection(changes)
        if approval_fields:
            raise HTTPException(
                status_code=409,
                detail=(
                    "Changes to department, degree, title, or employment status "
                    "must be submitted for approval."
                ),
            )
        for field, value in changes.items():
            setattr(teacher, field, value)
        record_audit(db, actor, "teacher_updated", "teacher", teacher.id, details={"fields": sorted(changes)})
        db.commit()
        db.refresh(teacher)
        return teacher

    @application.get(
        "/teachers/{teacher_id}/qualifications",
        response_model=list[QualificationRead],
        tags=["qualifications"],
    )
    def list_qualifications(teacher_id: int, db: DbSession, _: HR_USER) -> list[Qualification]:
        _require_teacher(db, teacher_id)
        return list(
            db.scalars(
                select(Qualification)
                .where(Qualification.teacher_id == teacher_id)
                .order_by(Qualification.issued_year.desc(), Qualification.id.desc())
            )
        )

    @application.post(
        "/teachers/{teacher_id}/qualifications",
        response_model=QualificationRead,
        status_code=status.HTTP_201_CREATED,
        tags=["qualifications"],
    )
    def create_qualification(
        teacher_id: int,
        payload: QualificationCreate,
        db: DbSession,
        actor: HR_USER,
    ) -> Qualification:
        _require_teacher(db, teacher_id)
        qualification = Qualification(teacher_id=teacher_id, **payload.model_dump())
        db.add(qualification)
        db.flush()
        record_audit(db, actor, "qualification_created", "qualification", qualification.id)
        db.commit()
        db.refresh(qualification)
        return qualification

    @application.patch(
        "/qualifications/{qualification_id}",
        response_model=QualificationRead,
        tags=["qualifications"],
    )
    def update_qualification(
        qualification_id: int,
        payload: QualificationUpdate,
        db: DbSession,
        actor: HR_USER,
    ) -> Qualification:
        qualification = db.get(Qualification, qualification_id)
        if qualification is None:
            raise HTTPException(status_code=404, detail="Qualification not found.")
        for field, value in payload.model_dump(exclude_unset=True).items():
            setattr(qualification, field, value)
        record_audit(db, actor, "qualification_updated", "qualification", qualification.id)
        db.commit()
        db.refresh(qualification)
        return qualification

    @application.get(
        "/teachers/{teacher_id}/contracts",
        response_model=list[ContractRead],
        tags=["contracts"],
    )
    def list_contracts(teacher_id: int, db: DbSession, _: HR_USER) -> list[dict[str, Any]]:
        _require_teacher(db, teacher_id)
        contracts = list(
            db.scalars(
                select(EmploymentContract)
                .where(EmploymentContract.teacher_id == teacher_id)
                .order_by(EmploymentContract.starts_on.desc(), EmploymentContract.id.desc())
            )
        )
        warning_days = _contract_warning_days(db)
        return [_contract_read_response(contract, warning_days) for contract in contracts]

    @application.post(
        "/teachers/{teacher_id}/contracts",
        response_model=ContractRead,
        status_code=status.HTTP_201_CREATED,
        tags=["contracts"],
    )
    def create_contract(
        teacher_id: int,
        payload: ContractCreate,
        db: DbSession,
        actor: HR_USER,
    ) -> dict[str, Any]:
        _require_teacher(db, teacher_id)
        contract = EmploymentContract(teacher_id=teacher_id, **payload.model_dump())
        db.add(contract)
        db.flush()
        record_audit(db, actor, "contract_created", "contract", contract.id)
        db.commit()
        db.refresh(contract)
        return _contract_read_response(contract, _contract_warning_days(db))

    @application.patch(
        "/contracts/{contract_id}",
        response_model=ContractRead,
        tags=["contracts"],
    )
    def update_contract(
        contract_id: int,
        payload: ContractUpdate,
        db: DbSession,
        actor: HR_USER,
    ) -> dict[str, Any]:
        contract = db.get(EmploymentContract, contract_id)
        if contract is None:
            raise HTTPException(status_code=404, detail="Contract not found.")
        changes = payload.model_dump(exclude_unset=True)
        starts_on = changes.get("starts_on", contract.starts_on)
        ends_on = changes.get("ends_on", contract.ends_on)
        if ends_on and ends_on < starts_on:
            raise HTTPException(status_code=422, detail="Contract end date cannot be before its start date.")
        for field, value in changes.items():
            setattr(contract, field, value)
        record_audit(db, actor, "contract_updated", "contract", contract.id)
        db.commit()
        db.refresh(contract)
        return _contract_read_response(contract, _contract_warning_days(db))

    @application.get(
        "/teachers/{teacher_id}/work-history",
        response_model=list[WorkHistoryRead],
        tags=["work-history"],
    )
    def list_work_history(
        teacher_id: int,
        db: DbSession,
        _: HR_USER,
    ) -> list[WorkHistory]:
        _require_teacher(db, teacher_id)
        return list(
            db.scalars(
                select(WorkHistory)
                .where(WorkHistory.teacher_id == teacher_id)
                .order_by(WorkHistory.starts_on, WorkHistory.id)
            )
        )

    @application.post(
        "/teachers/{teacher_id}/work-history",
        response_model=WorkHistoryRead,
        status_code=status.HTTP_201_CREATED,
        tags=["work-history"],
    )
    def create_work_history(
        teacher_id: int,
        payload: WorkHistoryCreate,
        db: DbSession,
        actor: HR_USER,
    ) -> WorkHistory:
        teacher = _require_teacher(db, teacher_id)
        if db.get(Department, payload.department_id) is None:
            raise HTTPException(status_code=404, detail="Department not found.")
        _validate_work_period(
            db, teacher_id, payload.starts_on, payload.ends_on
        )
        history = WorkHistory(teacher_id=teacher.id, **payload.model_dump())
        db.add(history)
        record_audit(db, actor, "work_history_created", "teacher", teacher.id)
        db.commit()
        db.refresh(history)
        return history

    @application.get(
        "/teachers/{teacher_id}/documents",
        response_model=list[EvidenceDocumentRead],
        tags=["documents"],
    )
    def list_documents(teacher_id: int, db: DbSession, _: HR_USER) -> list[EvidenceDocument]:
        _require_teacher(db, teacher_id)
        return list(
            db.scalars(
                select(EvidenceDocument)
                .where(EvidenceDocument.teacher_id == teacher_id)
                .order_by(EvidenceDocument.created_at.desc())
            )
        )

    @application.post(
        "/teachers/{teacher_id}/documents",
        response_model=EvidenceDocumentRead,
        status_code=status.HTTP_201_CREATED,
        tags=["documents"],
    )
    async def upload_document(
        teacher_id: int,
        db: DbSession,
        actor: HR_USER,
        file: UploadFile = File(...),
        checklist_item_id: int | None = Form(default=None),
        qualification_id: int | None = Form(default=None),
        contract_id: int | None = Form(default=None),
    ) -> EvidenceDocument:
        _require_teacher(db, teacher_id)
        if checklist_item_id is not None:
            item = db.get(ChecklistItem, checklist_item_id)
            if item is None or not item.is_active:
                raise HTTPException(status_code=404, detail="Active checklist item not found.")
        if qualification_id is not None:
            qualification = db.get(Qualification, qualification_id)
            if qualification is None or qualification.teacher_id != teacher_id:
                raise HTTPException(status_code=404, detail="Qualification not found.")
        if contract_id is not None:
            contract = db.get(EmploymentContract, contract_id)
            if contract is None or contract.teacher_id != teacher_id:
                raise HTTPException(status_code=404, detail="Contract not found.")
        if qualification_id is not None and contract_id is not None:
            raise HTTPException(status_code=422, detail="A document cannot belong to both a qualification and contract.")

        original_name = _safe_original_filename(file.filename or "evidence")
        extension = Path(original_name).suffix.lower()
        file_spec = ALLOWED_UPLOADS.get(extension)
        if file_spec is None:
            raise HTTPException(status_code=415, detail="Only PDF, PNG, and JPEG files are allowed.")
        content = await file.read(MAX_UPLOAD_BYTES + 1)
        if not content or len(content) > MAX_UPLOAD_BYTES:
            raise HTTPException(status_code=413, detail="File is empty or exceeds the 10 MB limit.")
        expected_type, validate_signature = file_spec
        if not validate_signature(content):
            raise HTTPException(status_code=415, detail="File content does not match its extension.")
        if file.content_type and file.content_type not in {expected_type, "application/octet-stream"}:
            raise HTTPException(status_code=415, detail="Unsupported content type.")

        stored_name = f"{secrets.token_hex(24)}{extension}"
        destination = application.state.upload_dir / stored_name
        with destination.open("xb") as output:
            output.write(content)
        document = EvidenceDocument(
            teacher_id=teacher_id,
            checklist_item_id=checklist_item_id,
            qualification_id=qualification_id,
            contract_id=contract_id,
            original_name=original_name,
            stored_name=stored_name,
            content_type=expected_type,
            size_bytes=len(content),
            uploaded_by=actor.id,
        )
        db.add(document)
        try:
            db.flush()
            record_audit(db, actor, "document_uploaded", "document", document.id)
            db.commit()
        except SQLAlchemyError:
            db.rollback()
            destination.unlink(missing_ok=True)
            raise
        db.refresh(document)
        return document

    @application.get(
        "/documents/{document_id}/download",
        tags=["documents"],
    )
    def download_document(
        document_id: int,
        db: DbSession,
        _: HR_USER,
    ) -> FileResponse:
        document = db.get(EvidenceDocument, document_id)
        if document is None:
            raise HTTPException(status_code=404, detail="Document not found.")
        path = application.state.upload_dir / document.stored_name
        if not path.is_file():
            raise HTTPException(status_code=404, detail="Stored document is missing.")
        return FileResponse(
            path,
            media_type=document.content_type,
            filename=document.original_name,
            content_disposition_type="attachment",
            headers={
                "Cache-Control": "private, no-store",
                "X-Content-Type-Options": "nosniff",
            },
        )

    @application.get(
        "/checklist-items",
        response_model=list[ChecklistItemRead],
        tags=["completeness"],
    )
    def list_checklist_items(db: DbSession, _: HR_USER) -> list[ChecklistItem]:
        return list(
            db.scalars(
                select(ChecklistItem).order_by(
                    ChecklistItem.category, ChecklistItem.label, ChecklistItem.id
                )
            )
        )

    @application.post(
        "/checklist-items",
        response_model=ChecklistItemRead,
        status_code=status.HTTP_201_CREATED,
        tags=["completeness"],
    )
    def create_checklist_item(
        payload: ChecklistItemCreate,
        db: DbSession,
        actor: HR_USER,
    ) -> ChecklistItem:
        item = ChecklistItem(**payload.model_dump())
        db.add(item)
        try:
            db.flush()
            for teacher in db.scalars(select(Teacher)):
                db.add(
                    TeacherChecklistRequirement(
                        teacher_id=teacher.id,
                        checklist_item_id=item.id,
                        status=_default_checklist_status(item),
                    )
                )
            record_audit(db, actor, "checklist_item_created", "checklist_item", item.id)
            db.commit()
        except IntegrityError as error:
            db.rollback()
            raise HTTPException(status_code=409, detail="Checklist code already exists.") from error
        db.refresh(item)
        return item

    @application.patch(
        "/checklist-items/{item_id}",
        response_model=ChecklistItemRead,
        tags=["completeness"],
    )
    def update_checklist_item(
        item_id: int,
        payload: ChecklistItemUpdate,
        db: DbSession,
        actor: HR_USER,
    ) -> ChecklistItem:
        item = db.get(ChecklistItem, item_id)
        if item is None:
            raise HTTPException(status_code=404, detail="Checklist item not found.")
        changes = payload.model_dump(exclude_unset=True)
        for field, value in changes.items():
            setattr(item, field, value)
        if changes.get("is_active") is True:
            existing_teacher_ids = set(
                db.scalars(
                    select(TeacherChecklistRequirement.teacher_id).where(
                        TeacherChecklistRequirement.checklist_item_id == item_id
                    )
                )
            )
            for teacher in db.scalars(select(Teacher)):
                if teacher.id not in existing_teacher_ids:
                    db.add(
                        TeacherChecklistRequirement(
                            teacher_id=teacher.id,
                            checklist_item_id=item.id,
                            status=_default_checklist_status(item),
                        )
                    )
        record_audit(db, actor, "checklist_item_updated", "checklist_item", item.id)
        db.commit()
        db.refresh(item)
        return item

    @application.get(
        "/teachers/{teacher_id}/checklist-requirements",
        response_model=list[TeacherChecklistRequirementRead],
        tags=["completeness"],
    )
    def list_teacher_checklist_requirements(
        teacher_id: int,
        db: DbSession,
        _: HR_USER,
    ) -> list[dict[str, Any]]:
        _require_teacher(db, teacher_id)
        rows = db.execute(
            select(ChecklistItem, TeacherChecklistRequirement)
            .outerjoin(
                TeacherChecklistRequirement,
                (TeacherChecklistRequirement.checklist_item_id == ChecklistItem.id)
                & (TeacherChecklistRequirement.teacher_id == teacher_id),
            )
            .where(ChecklistItem.is_active.is_(True))
            .order_by(ChecklistItem.category, ChecklistItem.label)
        ).all()
        return [
            {
                "item_id": item.id,
                "code": item.code,
                "label": item.label,
                "category": item.category,
                "classification": item.classification,
                "applicability": item.applicability,
                "is_sensitive": item.is_sensitive,
                "status": requirement.status if requirement else _default_checklist_status(item),
                "reviewed_at": requirement.reviewed_at if requirement else None,
            }
            for item, requirement in rows
        ]

    @application.patch(
        "/teachers/{teacher_id}/checklist-requirements/{item_id}",
        response_model=TeacherChecklistRequirementRead,
        tags=["completeness"],
    )
    def update_teacher_checklist_requirement(
        teacher_id: int,
        item_id: int,
        payload: TeacherChecklistRequirementUpdate,
        db: DbSession,
        actor: HR_USER,
    ) -> dict[str, Any]:
        _require_teacher(db, teacher_id)
        item = db.get(ChecklistItem, item_id)
        if item is None or not item.is_active:
            raise HTTPException(status_code=404, detail="Active checklist item not found.")
        requirement = db.scalar(
            select(TeacherChecklistRequirement).where(
                TeacherChecklistRequirement.teacher_id == teacher_id,
                TeacherChecklistRequirement.checklist_item_id == item_id,
            )
        )
        if requirement is None:
            requirement = TeacherChecklistRequirement(
                teacher_id=teacher_id,
                checklist_item_id=item_id,
                status=_default_checklist_status(item),
            )
            db.add(requirement)
        requirement.status = payload.status
        requirement.reviewed_by = actor.id if payload.status != "review" else None
        requirement.reviewed_at = datetime.now(timezone.utc) if payload.status != "review" else None
        record_audit(
            db,
            actor,
            "teacher_checklist_requirement_updated",
            "teacher_checklist_requirement",
            f"{teacher_id}:{item_id}",
        )
        db.commit()
        return {
            "item_id": item.id,
            "code": item.code,
            "label": item.label,
            "category": item.category,
            "classification": item.classification,
            "applicability": item.applicability,
            "is_sensitive": item.is_sensitive,
            "status": requirement.status,
            "reviewed_at": requirement.reviewed_at,
        }

    @application.get(
        "/teachers/{teacher_id}/completeness",
        tags=["completeness"],
    )
    def teacher_completeness(
        teacher_id: int,
        db: DbSession,
        _: HR_USER,
    ) -> dict[str, Any]:
        teacher = _require_teacher(db, teacher_id)
        missing, needs_review = _teacher_completeness(db, teacher)
        return {
            "teacher_id": teacher_id,
            "is_complete": not missing and not needs_review,
            "missing": missing,
            "needs_review": needs_review,
        }

    @application.get("/contracts/alerts", tags=["contracts"])
    def contract_alerts(
        db: DbSession,
        _: HR_USER,
        days: int | None = Query(default=None, ge=0, le=365),
    ) -> list[dict[str, Any]]:
        threshold = days if days is not None else _contract_warning_days(db)
        today = date.today()
        end_date = today + timedelta(days=threshold)
        contracts = db.execute(
            select(EmploymentContract, Teacher)
            .join(Teacher, EmploymentContract.teacher_id == Teacher.id)
            .where(
                EmploymentContract.status == "active",
                EmploymentContract.ends_on.is_not(None),
                EmploymentContract.ends_on <= end_date,
            )
            .order_by(EmploymentContract.ends_on)
        ).all()
        return [
            {
                "contract_id": contract.id,
                "teacher_id": teacher.id,
                "teacher_name": teacher.full_name,
                "contract_code": contract.contract_code,
                "ends_on": contract.ends_on,
                "days_remaining": (contract.ends_on - today).days,
                "status": "expired" if contract.ends_on < today else (
                    "expiring" if contract.ends_on <= end_date else "active"
                ),
            }
            for contract, teacher in contracts
        ]

    @application.get("/settings/contract-warning-days", tags=["settings"])
    def get_contract_warning_days(db: DbSession, _: HR_USER) -> dict[str, int]:
        return {"days": _contract_warning_days(db)}

    @application.put("/settings/contract-warning-days", tags=["settings"])
    def set_contract_warning_days(
        db: DbSession,
        actor: HR_USER,
        days: int = Query(30, ge=0, le=365),
    ) -> dict[str, int]:
        setting = db.get(AppSetting, "contract_warning_days")
        if setting is None:
            setting = AppSetting(key="contract_warning_days", value=str(days))
            db.add(setting)
        else:
            setting.value = str(days)
        record_audit(db, actor, "contract_warning_days_updated", "setting", "contract_warning_days")
        db.commit()
        return {"days": days}

    @application.get("/reports/summary", tags=["reports"])
    def report_summary(
        db: DbSession,
        _: HR_OR_MANAGER,
        department_id: int | None = Query(default=None, gt=0),
        start_date: date | None = None,
        end_date: date | None = None,
    ) -> dict[str, Any]:
        if start_date and end_date and end_date < start_date:
            raise HTTPException(status_code=422, detail="End date cannot be before start date.")
        statement = select(Teacher)
        if department_id is not None:
            statement = statement.where(Teacher.department_id == department_id)
        if start_date is not None:
            statement = statement.where(func.date(Teacher.created_at) >= start_date.isoformat())
        if end_date is not None:
            statement = statement.where(func.date(Teacher.created_at) <= end_date.isoformat())
        teachers = list(db.scalars(statement))
        teacher_ids = [teacher.id for teacher in teachers]
        complete_count = sum(not _missing_requirements(db, teacher) for teacher in teachers)
        status_counts: dict[str, int] = {}
        degree_counts: dict[str, int] = {}
        for teacher in teachers:
            status_counts[teacher.employment_status] = status_counts.get(teacher.employment_status, 0) + 1
            if teacher.academic_degree:
                degree_counts[teacher.academic_degree] = degree_counts.get(teacher.academic_degree, 0) + 1
        department_counts: dict[str, int] = {}
        for teacher in teachers:
            name = teacher.department.name if teacher.department else "Unknown"
            department_counts[name] = department_counts.get(name, 0) + 1

        today = date.today()
        warning_days = _contract_warning_days(db)
        contract_statement = select(EmploymentContract).where(
            EmploymentContract.status == "active",
            EmploymentContract.ends_on.is_not(None),
            EmploymentContract.ends_on <= today + timedelta(days=warning_days),
        )
        if department_id is not None:
            contract_statement = contract_statement.join(Teacher).where(
                Teacher.department_id == department_id
            )
        contracts = list(db.scalars(contract_statement))
        if teacher_ids:
            request_statement = select(ChangeRequest).where(
                ChangeRequest.teacher_id.in_(teacher_ids)
            )
        else:
            request_statement = select(ChangeRequest).where(ChangeRequest.id == -1)
        request_counts: dict[str, int] = {}
        for item in db.scalars(request_statement):
            request_counts[item.status] = request_counts.get(item.status, 0) + 1

        return {
            "filters": {
                "department_id": department_id,
                "start_date": start_date,
                "end_date": end_date,
                "date_basis": "teacher record creation date",
            },
            "total_teachers": len(teachers),
            "teachers_by_status": status_counts,
            "teachers_by_degree": degree_counts,
            "teachers_by_department": department_counts,
            "incomplete_profiles": len(teachers) - complete_count,
            "expiring_or_expired_contracts": len(contracts),
            "change_requests_by_status": request_counts,
            "contract_warning_days": warning_days,
        }

    @application.get("/reports/export.xlsx", tags=["reports"])
    def export_report(
        db: DbSession,
        actor: HR_OR_MANAGER,
        department_id: int | None = Query(default=None, gt=0),
        start_date: date | None = None,
        end_date: date | None = None,
    ) -> StreamingResponse:
        summary = report_summary(db, actor, department_id, start_date, end_date)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Giang vien"
        sheet.append(["Báo cáo hồ sơ giảng viên"])
        sheet.append(["Điều kiện lọc", json.dumps(summary["filters"], ensure_ascii=False, default=str)])
        if actor.role == "hr":
            sheet.append(["Mã giảng viên", "Họ tên", "Đơn vị", "Trình độ", "Chức danh", "Trạng thái", "Đầy đủ hồ sơ"])
        statement = select(Teacher)
        if department_id is not None:
            statement = statement.where(Teacher.department_id == department_id)
        if start_date is not None:
            statement = statement.where(func.date(Teacher.created_at) >= start_date.isoformat())
        if end_date is not None:
            statement = statement.where(func.date(Teacher.created_at) <= end_date.isoformat())
        if actor.role == "hr":
            for teacher in db.scalars(statement.order_by(Teacher.full_name)):
                sheet.append(
                    [
                        teacher.employee_code,
                        teacher.full_name,
                        teacher.department.name if teacher.department else "",
                        teacher.academic_degree or "",
                        teacher.academic_title or "",
                        teacher.employment_status,
                        "Đầy đủ" if not _missing_requirements(db, teacher) else "Còn thiếu",
                    ]
                )
        summary_sheet = workbook.create_sheet("Tong hop")
        summary_sheet.append(["Chỉ số", "Giá trị"])
        summary_sheet.append(["Tổng giảng viên", summary["total_teachers"]])
        summary_sheet.append(["Hồ sơ chưa đầy đủ", summary["incomplete_profiles"]])
        summary_sheet.append(["Hợp đồng sắp/đã hết hạn", summary["expiring_or_expired_contracts"]])
        for label, values in (
            ("Trạng thái", summary["teachers_by_status"]),
            ("Trình độ", summary["teachers_by_degree"]),
            ("Đơn vị", summary["teachers_by_department"]),
            ("Yêu cầu thay đổi", summary["change_requests_by_status"]),
        ):
            for value, count in values.items():
                summary_sheet.append([f"{label}: {value}", count])
        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="teacher_report.xlsx"'},
        )

    @application.post(
        "/teachers/{teacher_id}/change-requests",
        response_model=ChangeRequestRead,
        status_code=status.HTTP_201_CREATED,
        tags=["change-requests"],
    )
    def create_change_request(
        teacher_id: int,
        payload: ChangeRequestCreate,
        db: DbSession,
        actor: HR_USER,
    ) -> dict[str, Any]:
        teacher = _require_teacher(db, teacher_id)
        _validate_proposed_change(db, payload.proposed_value)
        current_value = {
            field: getattr(teacher, field) for field in payload.proposed_value
        }
        request_item = ChangeRequest(
            teacher_id=teacher.id,
            requester_id=actor.id,
            change_type=payload.change_type,
            current_value=json.dumps(current_value, ensure_ascii=False),
            proposed_value=json.dumps(payload.proposed_value, ensure_ascii=False),
            reason=payload.reason,
            effective_on=payload.effective_on,
        )
        db.add(request_item)
        db.flush()
        db.add(
            ChangeRequestEvent(
                change_request_id=request_item.id,
                actor_id=actor.id,
                event_type="submitted",
                note=payload.reason,
            )
        )
        record_audit(db, actor, "change_request_created", "change_request", request_item.id)
        db.commit()
        db.refresh(request_item)
        return _change_request_response(request_item)

    @application.get(
        "/change-requests",
        response_model=list[ChangeRequestRead],
        tags=["change-requests"],
    )
    def list_change_requests(
        db: DbSession,
        actor: Annotated[User, Depends(require_roles("hr", "manager"))],
        state_filter: str | None = Query(
            default=None,
            alias="status",
            pattern="^(pending|needs_info|approved|rejected)$",
        ),
    ) -> list[dict[str, Any]]:
        statement = select(ChangeRequest)
        if actor.role == "hr":
            statement = statement.where(ChangeRequest.requester_id == actor.id)
        elif state_filter is None:
            statement = statement.where(ChangeRequest.status == "pending")
        if state_filter:
            statement = statement.where(ChangeRequest.status == state_filter)
        return [
            _change_request_response(item)
            for item in db.scalars(statement.order_by(ChangeRequest.created_at.desc()))
        ]

    @application.patch(
        "/change-requests/{request_id}",
        response_model=ChangeRequestRead,
        tags=["change-requests"],
    )
    def update_change_request(
        request_id: int,
        payload: ChangeRequestUpdate,
        db: DbSession,
        actor: HR_USER,
    ) -> dict[str, Any]:
        item = db.get(ChangeRequest, request_id)
        if item is None:
            raise HTTPException(status_code=404, detail="Change request not found.")
        if item.requester_id != actor.id:
            raise HTTPException(status_code=403, detail="Only the requester can provide additional information.")
        if item.status != "needs_info":
            raise HTTPException(status_code=409, detail="This request is not awaiting additional information.")
        _validate_proposed_change(db, payload.proposed_value)
        teacher = _require_teacher(db, item.teacher_id)
        current_value = {
            field: getattr(teacher, field) for field in payload.proposed_value
        }
        item.proposed_value = json.dumps(payload.proposed_value, ensure_ascii=False)
        item.current_value = json.dumps(current_value, ensure_ascii=False)
        item.reason = payload.reason
        item.effective_on = payload.effective_on
        item.status = "pending"
        item.processor_id = None
        item.processed_at = None
        item.resolution_reason = None
        db.add(
            ChangeRequestEvent(
                change_request_id=item.id,
                actor_id=actor.id,
                event_type="resubmitted",
                note=payload.reason,
            )
        )
        record_audit(db, actor, "change_request_resubmitted", "change_request", item.id)
        db.commit()
        db.refresh(item)
        return _change_request_response(item)

    @application.post(
        "/change-requests/{request_id}/decision",
        response_model=ChangeRequestRead,
        tags=["change-requests"],
    )
    def decide_change_request(
        request_id: int,
        payload: ChangeDecision,
        db: DbSession,
        actor: MANAGER_USER,
    ) -> dict[str, Any]:
        item = db.get(ChangeRequest, request_id)
        if item is None:
            raise HTTPException(status_code=404, detail="Change request not found.")
        if item.status != "pending":
            raise HTTPException(status_code=409, detail="Only pending requests can be processed.")
        if item.requester_id == actor.id:
            raise HTTPException(status_code=403, detail="You cannot approve or process your own request.")
        claimed = db.execute(
            update(ChangeRequest)
            .where(ChangeRequest.id == request_id, ChangeRequest.status == "pending")
            .values(status="processing")
        )
        if claimed.rowcount != 1:
            db.rollback()
            raise HTTPException(status_code=409, detail="This request has already been processed.")
        item.status = "processing"
        item.processor_id = actor.id
        item.processed_at = datetime.now(timezone.utc)
        item.resolution_reason = payload.reason
        if payload.action == "approve":
            _apply_approved_change(db, item)
            item.status = "approved"
        elif payload.action == "reject":
            item.status = "rejected"
        else:
            item.status = "needs_info"
        db.add(
            ChangeRequestEvent(
                change_request_id=item.id,
                actor_id=actor.id,
                event_type=payload.action,
                note=payload.reason,
            )
        )
        record_audit(
            db,
            actor,
            f"change_request_{payload.action}",
            "change_request",
            item.id,
            details={"teacher_id": item.teacher_id},
        )
        db.commit()
        db.refresh(item)
        return _change_request_response(item)

    @application.get(
        "/change-requests/{request_id}/events",
        tags=["change-requests"],
    )
    def list_change_request_events(
        request_id: int,
        db: DbSession,
        actor: Annotated[User, Depends(require_roles("hr", "manager"))],
    ) -> list[dict[str, Any]]:
        item = db.get(ChangeRequest, request_id)
        if item is None:
            raise HTTPException(status_code=404, detail="Change request not found.")
        if actor.role == "hr" and item.requester_id != actor.id:
            raise HTTPException(status_code=403, detail="You can only view your own requests.")
        events = db.scalars(
            select(ChangeRequestEvent)
            .where(ChangeRequestEvent.change_request_id == request_id)
            .order_by(ChangeRequestEvent.created_at, ChangeRequestEvent.id)
        )
        return [
            {
                "id": event.id,
                "actor_id": event.actor_id,
                "event_type": event.event_type,
                "note": event.note,
                "created_at": event.created_at,
            }
            for event in events
        ]

    @application.get("/audit-logs", tags=["audit"])
    def list_audit_logs(
        db: DbSession,
        _: ADMIN_USER,
        actor_id: int | None = Query(default=None, gt=0),
        action: str | None = Query(default=None, min_length=1, max_length=80),
        start_date: date | None = None,
        end_date: date | None = None,
        limit: int = Query(default=100, ge=1, le=500),
    ) -> list[dict[str, Any]]:
        statement = select(AuditLog)
        if actor_id is not None:
            statement = statement.where(AuditLog.actor_id == actor_id)
        if action:
            statement = statement.where(AuditLog.action == action)
        if start_date:
            statement = statement.where(func.date(AuditLog.created_at) >= start_date.isoformat())
        if end_date:
            statement = statement.where(func.date(AuditLog.created_at) <= end_date.isoformat())
        logs = db.scalars(statement.order_by(AuditLog.created_at.desc()).limit(limit))
        return [
            {
                "id": item.id,
                "actor_id": item.actor_id,
                "action": item.action,
                "entity_type": item.entity_type,
                "entity_id": item.entity_id,
                "outcome": item.outcome,
                "details": json.loads(item.details) if item.details else None,
                "created_at": item.created_at,
            }
            for item in logs
        ]

    return application


def _require_csrf(request: Request) -> None:
    login_session: LoginSession = getattr(request.state, "login_session", None)
    supplied = request.headers.get("X-CSRF-Token", "")
    if login_session is None or not secrets.compare_digest(
        hash_secret(supplied), login_session.csrf_token_hash
    ):
        raise HTTPException(status_code=403, detail="A valid CSRF token is required.")


def _as_utc(value: datetime) -> datetime:
    return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)


def _validate_parent_department(
    db: Session,
    parent_id: int | None,
    department_id: int | None = None,
) -> None:
    if parent_id is None:
        return
    parent = db.get(Department, parent_id)
    if parent is None:
        raise HTTPException(status_code=404, detail="Parent department not found.")
    if department_id == parent_id:
        raise HTTPException(status_code=422, detail="A department cannot be its own parent.")
    visited = {department_id} if department_id is not None else set()
    current: Department | None = parent
    while current is not None:
        if current.id in visited:
            raise HTTPException(status_code=422, detail="Department hierarchy cannot contain a cycle.")
        visited.add(current.id)
        current = current.parent


def _require_teacher(db: Session, teacher_id: int) -> Teacher:
    teacher = db.get(Teacher, teacher_id)
    if teacher is None:
        raise HTTPException(status_code=404, detail="Teacher not found.")
    return teacher


def _safe_original_filename(filename: str) -> str:
    normalized = filename.replace("\\", "/").split("/")[-1].strip()
    if not normalized or normalized in {".", ".."}:
        return "evidence"
    normalized = "".join("_" if ord(character) < 32 else character for character in normalized)
    return normalized[:255]


def _default_checklist_status(item: ChecklistItem) -> str:
    return "optional" if item.classification == "optional" else "review"


def _teacher_completeness(
    db: Session,
    teacher: Teacher,
) -> tuple[list[str], list[str]]:
    missing: list[str] = []
    needs_review: list[str] = []
    if not teacher.employee_code:
        missing.append("Mã giảng viên")
    if not teacher.full_name:
        missing.append("Họ và tên")
    if teacher.department_id is None:
        missing.append("Đơn vị công tác")
    requirements = db.execute(
        select(ChecklistItem, TeacherChecklistRequirement.status)
        .outerjoin(
            TeacherChecklistRequirement,
            (TeacherChecklistRequirement.checklist_item_id == ChecklistItem.id)
            & (TeacherChecklistRequirement.teacher_id == teacher.id),
        )
        .where(ChecklistItem.is_active.is_(True))
    )
    uploaded_item_ids = set(
        db.scalars(
            select(EvidenceDocument.checklist_item_id).where(
                EvidenceDocument.teacher_id == teacher.id,
                EvidenceDocument.checklist_item_id.is_not(None),
            )
        )
    )
    for item, status in requirements:
        requirement_status = status or _default_checklist_status(item)
        if requirement_status == "review":
            needs_review.append(item.label)
        elif requirement_status == "required" and item.id not in uploaded_item_ids:
            missing.append(item.label)
    return missing, needs_review


def _missing_requirements(db: Session, teacher: Teacher) -> list[str]:
    missing, needs_review = _teacher_completeness(db, teacher)
    return missing + [f"Cần rà soát điều kiện: {label}" for label in needs_review]


def _contract_warning_days(db: Session) -> int:
    setting = db.get(AppSetting, "contract_warning_days")
    if setting is None:
        return 30
    try:
        return min(max(int(setting.value), 0), 365)
    except ValueError:
        return 30


def _contract_read_response(
    contract: EmploymentContract,
    warning_days: int,
) -> dict[str, Any]:
    today = date.today()
    display_status = contract.status
    if contract.status == "active" and contract.ends_on is not None:
        if contract.ends_on < today:
            display_status = "expired"
        elif contract.ends_on <= today + timedelta(days=warning_days):
            display_status = "expiring"
    return {
        "id": contract.id,
        "teacher_id": contract.teacher_id,
        "contract_code": contract.contract_code,
        "contract_type": contract.contract_type,
        "signed_on": contract.signed_on,
        "starts_on": contract.starts_on,
        "ends_on": contract.ends_on,
        "status": contract.status,
        "display_status": display_status,
    }


def _validate_proposed_change(db: Session, proposed: dict[str, object]) -> None:
    unknown = set(proposed) - APPROVABLE_FIELDS
    if unknown:
        raise HTTPException(
            status_code=422,
            detail=f"Unsupported approved fields: {', '.join(sorted(unknown))}.",
        )
    if not proposed:
        raise HTTPException(status_code=422, detail="At least one approved field is required.")
    if "employment_status" in proposed and proposed["employment_status"] not in {
        "active",
        "paused",
        "ended",
    }:
        raise HTTPException(status_code=422, detail="Invalid employment status.")
    if "department_id" in proposed:
        department_id = proposed["department_id"]
        if type(department_id) is not int or db.get(Department, department_id) is None:
            raise HTTPException(status_code=422, detail="Proposed department does not exist.")
    for field in ("academic_degree", "academic_title"):
        if field in proposed and (
            proposed[field] is not None
            and (not isinstance(proposed[field], str) or len(proposed[field]) > 80)
        ):
            raise HTTPException(status_code=422, detail=f"Invalid value for {field}.")


def _apply_approved_change(db: Session, item: ChangeRequest) -> None:
    teacher = _require_teacher(db, item.teacher_id)
    proposed = json.loads(item.proposed_value)
    _validate_proposed_change(db, proposed)
    old_department_id = teacher.department_id
    old_title = teacher.academic_title
    effective_on = item.effective_on or date.today()
    if "department_id" in proposed or "academic_title" in proposed:
        current_history = db.scalar(
            select(WorkHistory)
            .where(
                WorkHistory.teacher_id == teacher.id,
                WorkHistory.ends_on.is_(None),
            )
            .order_by(WorkHistory.starts_on.desc(), WorkHistory.id.desc())
        )
        if current_history and effective_on <= current_history.starts_on:
            raise HTTPException(
                status_code=422,
                detail="Effective date must be after the current work period start.",
            )
        _validate_work_period(
            db,
            teacher.id,
            effective_on,
            None,
            exclude_id=current_history.id if current_history else None,
        )
        if current_history:
            current_history.ends_on = effective_on - timedelta(days=1)
        next_department_id = int(proposed.get("department_id", old_department_id))
        next_title = proposed.get("academic_title", old_title)
        if "department_id" in proposed:
            teacher.department_id = next_department_id
        if "academic_title" in proposed:
            teacher.academic_title = next_title
        db.add(
            WorkHistory(
                teacher_id=teacher.id,
                department_id=next_department_id,
                position_title=next_title,
                starts_on=effective_on,
                reason=item.reason,
                change_request_id=item.id,
            )
        )
    for field, value in proposed.items():
        setattr(teacher, field, value)


def _validate_work_period(
    db: Session,
    teacher_id: int,
    starts_on: date,
    ends_on: date | None,
    exclude_id: int | None = None,
) -> None:
    statement = select(WorkHistory).where(WorkHistory.teacher_id == teacher_id)
    if exclude_id is not None:
        statement = statement.where(WorkHistory.id != exclude_id)
    history = db.scalars(statement)
    for existing in history:
        start = existing.starts_on
        end = existing.ends_on
        overlaps = (ends_on is None or start <= ends_on) and (
            end is None or starts_on <= end
        )
        if overlaps:
            raise HTTPException(
                status_code=409,
                detail="Work history periods cannot overlap.",
            )


def _change_request_response(item: ChangeRequest) -> dict[str, Any]:
    return {
        "id": item.id,
        "teacher_id": item.teacher_id,
        "requester_id": item.requester_id,
        "change_type": item.change_type,
        "current_value": json.loads(item.current_value),
        "proposed_value": json.loads(item.proposed_value),
        "reason": item.reason,
        "effective_on": item.effective_on,
        "status": item.status,
        "processor_id": item.processor_id,
        "processed_at": item.processed_at,
        "resolution_reason": item.resolution_reason,
        "created_at": item.created_at,
    }


def app_factory() -> FastAPI:
    return create_app()


app = create_app()
