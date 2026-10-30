import tempfile
import unittest
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import create_app
from app.models import User
from app.security import hash_password


class TeacherManagementApiTests(unittest.TestCase):
    def setUp(self) -> None:
        self.iterations_patch = patch("app.security.PASSWORD_ITERATIONS", 1_000)
        self.iterations_patch.start()
        self.temporary_uploads = tempfile.TemporaryDirectory()
        self.client_context = TestClient(create_app("sqlite://"))
        self.client = self.client_context.__enter__()
        self.client.app.state.upload_dir = Path(self.temporary_uploads.name)
        password_hash = hash_password("Correct horse battery staple 2048!")
        with self.client.app.state.session_factory() as db:
            for username, role in (
                ("admin", "admin"),
                ("hr", "hr"),
                ("manager", "manager"),
            ):
                db.add(
                    User(
                        username=username,
                        full_name=username.title(),
                        password_hash=password_hash,
                        role=role,
                    )
                )
            db.commit()
        self.hr = self.login("hr")
        self.csrf = self.hr["csrf_token"]
        response = self.client.post(
            "/departments",
            json={"code": "IT", "name": "Information Technology"},
            headers=self.headers(),
        )
        self.assertEqual(response.status_code, 201, response.text)
        self.department_id = response.json()["id"]

    def tearDown(self) -> None:
        self.client_context.__exit__(None, None, None)
        self.temporary_uploads.cleanup()
        self.iterations_patch.stop()

    def login(
        self,
        username: str,
        password: str = "Correct horse battery staple 2048!",
    ) -> dict[str, object]:
        response = self.client.post(
            "/auth/login",
            json={"username": username, "password": password},
        )
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def headers(self) -> dict[str, str]:
        return {"X-CSRF-Token": str(self.csrf)}

    def create_teacher(self, employee_code: str = "T001", **overrides: object):
        payload: dict[str, object] = {
            "employee_code": employee_code,
            "full_name": "Nguyen Van An",
            "department_id": self.department_id,
            "academic_degree": "Master",
        }
        payload.update(overrides)
        return self.client.post("/teachers", json=payload, headers=self.headers())

    def test_authentication_static_assets_and_api_access(self) -> None:
        page = self.client.get("/")
        self.assertEqual(page.status_code, 200)
        self.assertIn("TeacherHub", page.text)
        self.assertIn('id="teacher-form"', page.text)
        self.assertEqual(self.client.get("/health").json(), {"status": "ok"})
        self.assertEqual(self.client.get("/static/app.js").status_code, 200)
        self.assertEqual(self.client.get("/static/styles.css").status_code, 200)
        self.assertEqual(self.client.get("/static/favicon.svg").status_code, 200)
        session = self.client.get("/auth/session")
        self.assertEqual(session.status_code, 200)
        self.assertEqual(session.json()["csrf_token"], self.csrf)
        self.assertEqual(page.headers["x-content-type-options"], "nosniff")

        unauthenticated = TestClient(create_app("sqlite://"))
        try:
            self.assertEqual(unauthenticated.get("/teachers").status_code, 401)
        finally:
            unauthenticated.close()

    def test_create_filter_and_duplicate_teacher_code(self) -> None:
        created = self.create_teacher()
        self.assertEqual(created.status_code, 201, created.text)
        self.assertEqual(created.json()["employee_code"], "T001")
        self.assertEqual(
            self.create_teacher("T002", full_name="Other Person", email="").json()["email"],
            None,
        )
        invalid_email = self.create_teacher("T003", email="not-an-email")
        self.assertEqual(invalid_email.status_code, 422)
        invalid_email_update = self.client.patch(
            f"/teachers/{created.json()['id']}",
            json={"email": "not-an-email"},
            headers=self.headers(),
        )
        self.assertEqual(invalid_email_update.status_code, 422)
        response = self.client.get(
            "/teachers",
            params={"q": "Nguyen", "academic_degree": "Master"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.json()), 1)

        duplicate = self.create_teacher()
        self.assertEqual(duplicate.status_code, 409)

    def test_backend_role_enforcement_and_csrf(self) -> None:
        invalid_csrf = self.client.post(
            "/departments",
            json={"code": "OPS", "name": "Operations"},
            headers={"X-CSRF-Token": "invalid"},
        )
        self.assertEqual(invalid_csrf.status_code, 403)

        admin = self.login("admin")
        response = self.client.post(
            "/users",
            json={
                "username": "new-user",
                "full_name": "New User",
                "password": "Another safe password 456!",
                "role": "hr",
            },
            headers={"X-CSRF-Token": admin["csrf_token"]},
        )
        self.assertEqual(response.status_code, 201, response.text)
        self.assertEqual(
            self.client.post(
                "/departments",
                json={"code": "HR", "name": "Human Resources"},
                headers={"X-CSRF-Token": admin["csrf_token"]},
            ).status_code,
            403,
        )
        manager = self.login("manager")
        self.assertEqual(self.client.get("/teachers").status_code, 403)
        self.assertEqual(
            self.client.get("/teachers/1/checklist-requirements").status_code,
            403,
        )
        self.assertEqual(
            self.client.get(
                "/reports/export.xlsx",
                headers={"X-CSRF-Token": manager["csrf_token"]},
            ).status_code,
            200,
        )

    def test_department_hierarchy_and_deactivation(self) -> None:
        child = self.client.post(
            "/departments",
            json={"code": "SYS", "name": "Systems", "parent_id": self.department_id},
            headers=self.headers(),
        )
        self.assertEqual(child.status_code, 201, child.text)
        self.assertEqual(child.json()["parent_id"], self.department_id)

        cycle = self.client.patch(
            f"/departments/{self.department_id}",
            json={"parent_id": child.json()["id"]},
            headers=self.headers(),
        )
        self.assertEqual(cycle.status_code, 422)

        response = self.client.patch(
            f"/departments/{self.department_id}",
            json={"is_active": False},
            headers=self.headers(),
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            self.client.get("/departments").json()[0]["id"],
            child.json()["id"],
        )

    def test_important_changes_are_approved_and_preserve_history(self) -> None:
        created = self.create_teacher()
        teacher_id = created.json()["id"]
        new_department = self.client.post(
            "/departments",
            json={"code": "LAW", "name": "Law"},
            headers=self.headers(),
        ).json()

        direct_change = self.client.patch(
            f"/teachers/{teacher_id}",
            json={"department_id": new_department["id"]},
            headers=self.headers(),
        )
        self.assertEqual(direct_change.status_code, 409)

        effective_on = date.today() + timedelta(days=1)
        request_response = self.client.post(
            f"/teachers/{teacher_id}/change-requests",
            json={
                "change_type": "department_transfer",
                "current_value": {"department_id": 999},
                "proposed_value": {"department_id": new_department["id"]},
                "reason": "Approved organizational transfer",
                "effective_on": effective_on.isoformat(),
            },
            headers=self.headers(),
        )
        self.assertEqual(request_response.status_code, 201, request_response.text)
        request_id = request_response.json()["id"]
        self.assertEqual(
            request_response.json()["current_value"]["department_id"],
            self.department_id,
        )

        own_decision = self.client.post(
            f"/change-requests/{request_id}/decision",
            json={"action": "approve"},
            headers=self.headers(),
        )
        self.assertEqual(own_decision.status_code, 403)

        manager = self.login("manager")
        decision = self.client.post(
            f"/change-requests/{request_id}/decision",
            json={"action": "approve"},
            headers={"X-CSRF-Token": manager["csrf_token"]},
        )
        self.assertEqual(decision.status_code, 200, decision.text)
        self.assertEqual(decision.json()["status"], "approved")

        self.login("hr")
        teacher = self.client.get(f"/teachers/{teacher_id}")
        self.assertEqual(teacher.json()["department_id"], new_department["id"])
        history = self.client.get(f"/teachers/{teacher_id}/work-history")
        self.assertEqual(len(history.json()), 2)
        self.assertEqual(
            history.json()[0]["ends_on"],
            (effective_on - timedelta(days=1)).isoformat(),
        )
        self.assertEqual(history.json()[1]["starts_on"], effective_on.isoformat())
        manager = self.login("manager")
        repeated = self.client.post(
            f"/change-requests/{request_id}/decision",
            json={"action": "approve"},
            headers={"X-CSRF-Token": manager["csrf_token"]},
        )
        self.assertEqual(repeated.status_code, 409)

    def test_contract_alerts_checklist_and_protected_upload(self) -> None:
        teacher_id = self.create_teacher().json()["id"]
        checklist = self.client.post(
            "/checklist-items",
            json={
                "code": "degree-proof",
                "label": "Bằng cấp minh chứng",
                "category": "qualification",
                "classification": "required",
                "applicability": "Theo yêu cầu vị trí.",
            },
            headers=self.headers(),
        )
        self.assertEqual(checklist.status_code, 201, checklist.text)
        requirements = self.client.get(
            f"/teachers/{teacher_id}/checklist-requirements"
        )
        self.assertEqual(requirements.status_code, 200)
        self.assertEqual(requirements.json()[0]["status"], "review")
        incomplete = self.client.get(f"/teachers/{teacher_id}/completeness").json()
        self.assertFalse(incomplete["is_complete"])
        self.assertIn("Bằng cấp minh chứng", incomplete["needs_review"])
        requirement = self.client.patch(
            f"/teachers/{teacher_id}/checklist-requirements/{checklist.json()['id']}",
            json={"status": "not_applicable"},
            headers=self.headers(),
        )
        self.assertEqual(requirement.status_code, 200, requirement.text)
        self.assertTrue(
            self.client.get(f"/teachers/{teacher_id}/completeness").json()["is_complete"]
        )
        requirement = self.client.patch(
            f"/teachers/{teacher_id}/checklist-requirements/{checklist.json()['id']}",
            json={"status": "required"},
            headers=self.headers(),
        )
        self.assertEqual(requirement.status_code, 200, requirement.text)
        incomplete = self.client.get(f"/teachers/{teacher_id}/completeness").json()
        self.assertIn("Bằng cấp minh chứng", incomplete["missing"])

        upload = self.client.post(
            f"/teachers/{teacher_id}/documents",
            data={"checklist_item_id": checklist.json()["id"]},
            files={"file": ("degree.pdf", b"%PDF-1.4\nproof", "application/pdf")},
            headers=self.headers(),
        )
        self.assertEqual(upload.status_code, 201, upload.text)
        document_id = upload.json()["id"]
        self.assertEqual(
            self.client.get(f"/teachers/{teacher_id}/completeness").json()["is_complete"],
            True,
        )
        downloaded = self.client.get(f"/documents/{document_id}/download")
        self.assertEqual(downloaded.status_code, 200)
        self.assertEqual(downloaded.content, b"%PDF-1.4\nproof")
        self.assertEqual(downloaded.headers["cache-control"], "private, no-store")

        contract = self.client.post(
            f"/teachers/{teacher_id}/contracts",
            json={
                "contract_code": "HD-001",
                "contract_type": "Fixed term",
                "starts_on": date.today().isoformat(),
                "ends_on": (date.today() + timedelta(days=10)).isoformat(),
            },
            headers=self.headers(),
        )
        self.assertEqual(contract.status_code, 201, contract.text)
        alerts = self.client.get("/contracts/alerts")
        self.assertEqual(alerts.status_code, 200)
        self.assertEqual(alerts.json()[0]["status"], "expiring")
        contract_list = self.client.get(f"/teachers/{teacher_id}/contracts")
        self.assertEqual(contract_list.json()[0]["display_status"], "expiring")

        invalid_contract = self.client.post(
            f"/teachers/{teacher_id}/contracts",
            json={
                "contract_code": "HD-002",
                "contract_type": "Invalid",
                "starts_on": date.today().isoformat(),
                "ends_on": (date.today() - timedelta(days=1)).isoformat(),
            },
            headers=self.headers(),
        )
        self.assertEqual(invalid_contract.status_code, 422)

    def test_reports_export_and_audit_access(self) -> None:
        self.create_teacher()
        summary = self.client.get("/reports/summary")
        self.assertEqual(summary.status_code, 200, summary.text)
        self.assertEqual(summary.json()["total_teachers"], 1)

        exported = self.client.get("/reports/export.xlsx")
        self.assertEqual(exported.status_code, 200)
        workbook = load_workbook(BytesIO(exported.content), read_only=True)
        self.assertIn("Giang vien", workbook.sheetnames)
        workbook.close()

        manager = self.login("manager")
        manager_export = self.client.get(
            "/reports/export.xlsx",
            headers={"X-CSRF-Token": manager["csrf_token"]},
        )
        self.assertEqual(manager_export.status_code, 200)
        manager_workbook = load_workbook(BytesIO(manager_export.content), read_only=True)
        manager_values = [
            cell.value
            for sheet in manager_workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
        ]
        self.assertNotIn("T001", manager_values)
        self.assertNotIn("Nguyen Van An", manager_values)
        manager_workbook.close()
        self.assertEqual(self.client.get("/audit-logs").status_code, 403)

        admin = self.login("admin")
        self.assertEqual(
            self.client.get(
                "/audit-logs",
                headers={"X-CSRF-Token": admin["csrf_token"]},
            ).status_code,
            200,
        )
        self.assertEqual(
            self.client.delete(
                "/audit-logs/1",
                headers={"X-CSRF-Token": admin["csrf_token"]},
            ).status_code,
            404,
        )

    def test_password_change_invalidates_session(self) -> None:
        response = self.client.post(
            "/auth/password",
            json={
                "current_password": "Correct horse battery staple 2048!",
                "new_password": "A different secure password 2049!",
            },
            headers=self.headers(),
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(self.client.get("/teachers").status_code, 401)
        new_login = self.login("hr", "A different secure password 2049!")
        self.assertTrue(new_login["csrf_token"])


if __name__ == "__main__":
    unittest.main()
